from __future__ import annotations

from pathlib import Path

import docx
import pymupdf
import pytest

from app.ingestion.parser import UnsupportedDocumentError, parse_document


def write_pdf(path: Path, pages: list[list[str]]) -> None:
    document = pymupdf.open()
    for blocks in pages:
        page = document.new_page()
        y = 72
        for text in blocks:
            page.insert_text((72, y), text, fontsize=12)
            y += 28
    document.save(str(path))
    document.close()


def test_markdown_is_read_as_markdown(tmp_path: Path) -> None:
    path = tmp_path / "note.md"
    path.write_text("# 标题\n正文。", encoding="utf-8")

    parsed = parse_document(path)

    assert (parsed.mode, parsed.doc_type) == ("markdown", "markdown")
    assert parsed.page_starts == ()
    assert parsed.page_count is None


def test_plain_text_uses_text_mode(tmp_path: Path) -> None:
    path = tmp_path / "note.txt"
    path.write_text("纯文本。", encoding="utf-8")

    assert parse_document(path).mode == "text"


def test_code_whitelist_uses_code_mode(tmp_path: Path) -> None:
    path = tmp_path / "chunker.py"
    path.write_text("def f():\n    return 1\n", encoding="utf-8")

    parsed = parse_document(path)

    assert (parsed.mode, parsed.doc_type) == ("code", "code")


def test_unknown_suffix_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "image.png"
    path.write_bytes(b"png")

    with pytest.raises(UnsupportedDocumentError, match="unsupported"):
        parse_document(path)


def test_legacy_doc_gets_an_actionable_message(tmp_path: Path) -> None:
    path = tmp_path / "old.doc"
    path.write_bytes(b"doc")

    with pytest.raises(UnsupportedDocumentError, match="docx"):
        parse_document(path)


def test_pdf_text_carries_page_starts_and_count(tmp_path: Path) -> None:
    path = tmp_path / "manual.pdf"
    write_pdf(
        path,
        [
            ["Chapter one introduces hybrid retrieval."],
            ["Chapter two covers reranking."],
        ],
    )

    parsed = parse_document(path)

    assert (parsed.mode, parsed.doc_type) == ("text", "pdf")
    assert parsed.page_count == 2
    assert len(parsed.page_starts) == 2
    assert parsed.page_starts[0] == 0
    # The offset points at the first character of page two.
    assert parsed.text[parsed.page_starts[1] :].startswith("Chapter two")


def test_scanned_pdf_is_reported_instead_of_indexed_empty(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    write_pdf(path, [[], []])  # pages with no text layer

    with pytest.raises(UnsupportedDocumentError, match="OCR"):
        parse_document(path)


def test_docx_headings_become_markdown_headings(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    document = docx.Document()
    document.add_heading("部署指南", level=1)
    document.add_paragraph("先备份数据库。")
    document.add_heading("回滚", level=2)
    document.add_paragraph("重新部署上一个版本。")
    document.save(str(path))

    parsed = parse_document(path)

    assert (parsed.mode, parsed.doc_type) == ("markdown", "docx")
    assert "# 部署指南" in parsed.text
    assert "## 回滚" in parsed.text
    assert parsed.page_starts == ()


def test_docx_tables_are_kept_as_rows(tmp_path: Path) -> None:
    path = tmp_path / "table.docx"
    document = docx.Document()
    document.add_paragraph("配置项如下。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "变量"
    table.cell(0, 1).text = "默认值"
    table.cell(1, 0).text = "RRF_K"
    table.cell(1, 1).text = "60"
    document.save(str(path))

    assert "变量 | 默认值" in parse_document(path).text


def test_empty_docx_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "blank.docx"
    docx.Document().save(str(path))

    with pytest.raises(UnsupportedDocumentError, match="没有可提取文本"):
        parse_document(path)
