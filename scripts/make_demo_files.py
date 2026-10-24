"""Generate the PDF and docx samples under examples/formats/.

The main demo corpus (examples/corpus) stays markdown-only so the 40-question
retrieval evaluation in eval/stage2 keeps measuring the same thing. These extra
files exist to show the stage-5 parsers and page-level citations in the UI.

Run it only when you want to regenerate the samples:

    .venv/Scripts/python.exe scripts/make_demo_files.py

Everything is generated locally; nothing is downloaded.
"""

from __future__ import annotations

from pathlib import Path

import docx
import pymupdf

OUT_DIR = Path(__file__).resolve().parents[1] / "examples" / "formats"

PDF_PAGES: list[list[str]] = [
    [
        "LocalMind 运维手册 · 第 1 页",
        "",
        "一、备份范围",
        "需要备份两样东西：data/localmind.db 与 Qdrant 的 volume。",
        "前者存资料源、文档、分块、会话与记忆；后者存稠密与稀疏向量。",
        "只备份数据库会丢掉向量，重建索引需要重新嵌入全部文件。",
        "",
        "二、备份频率",
        "建议每周一次，并在批量导入新资料后立即再备份一次。",
    ],
    [
        "LocalMind 运维手册 · 第 2 页",
        "",
        "三、端口约定",
        "后端只监听 127.0.0.1:8000，前端开发服务器 127.0.0.1:5173，",
        "Qdrant 使用 6333（HTTP）与 6334（gRPC），Ollama 使用 11434。",
        "所有端口都绑定环回地址，不对局域网开放。",
        "",
        "四、升级步骤",
        "先停后端，再拉取代码并执行 upgrade head，最后重启服务。",
        "数据库迁移在启动时执行，失败会直接终止启动而不是带着旧结构继续跑。",
    ],
    [
        "LocalMind 运维手册 · 第 3 页",
        "",
        "五、常见故障",
        "Qdrant 连不上：检查容器是否在运行，端口是否被占用。",
        "重排模型加载失败：确认 HF_HOME 指向的缓存目录存在且完整。",
        "PDF 索引后没有内容：多半是扫描版，没有文本层，需要 OCR 或换一份文件。",
        "",
        "六、日志位置",
        "服务日志输出到标准输出，每行一个 JSON，包含 event 字段便于过滤。",
    ],
]

DOCX_SECTIONS: list[tuple[int, str, list[str]]] = [
    (1, "知识库交付清单", []),
    (
        2,
        "一、资料入库",
        [
            "确认资料源目录只包含需要被检索的文件，避免把下载目录整目录加入。",
            "首次入库后检查分块数量与失败文件列表。",
        ],
    ),
    (
        2,
        "二、检索验收",
        [
            "用评测集跑一次检索指标，确认 Recall@5 不低于 0.8。",
            "抽查无答案问题，确认全部拒答且没有编造。",
        ],
    ),
    (
        2,
        "三、上线前检查",
        [
            "确认 .env 未入库，密钥只存在于本机。",
            "确认云端轨道按资料源的黑名单符合预期。",
            "确认本地轨道在断网状态下仍能完成一次问答。",
        ],
    ),
]

DOCX_TABLE = [
    ["检查项", "命令", "期望结果"],
    ["后端测试", "pytest backend/tests", "全部通过"],
    ["检索评测", "scripts/eval_retrieval.py", "Recall@5 >= 0.8"],
    ["离线验证", "scripts/verify_offline.py", "非环回连接：无"],
]


def make_pdf(path: Path) -> None:
    document = pymupdf.open()
    for lines in PDF_PAGES:
        page = document.new_page()
        y = 90
        for line in lines:
            if line:
                # china-s is a built-in CJK font, so no font file is needed.
                page.insert_text((72, y), line, fontname="china-s", fontsize=12)
            y += 26
    document.save(str(path))
    document.close()


def make_docx(path: Path) -> None:
    document = docx.Document()
    for level, title, paragraphs in DOCX_SECTIONS:
        document.add_heading(title, level=level)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    document.add_heading("四、命令速查", level=2)
    table = document.add_table(rows=len(DOCX_TABLE), cols=3)
    for row_index, row in enumerate(DOCX_TABLE):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    document.save(str(path))


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    pdf_path = OUT_DIR / "localmind-runbook.pdf"
    docx_path = OUT_DIR / "delivery-checklist.docx"
    make_pdf(pdf_path)
    make_docx(docx_path)
    print(f"wrote {pdf_path.relative_to(OUT_DIR.parents[1])} ({pdf_path.stat().st_size} bytes)")
    print(f"wrote {docx_path.relative_to(OUT_DIR.parents[1])} ({docx_path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
